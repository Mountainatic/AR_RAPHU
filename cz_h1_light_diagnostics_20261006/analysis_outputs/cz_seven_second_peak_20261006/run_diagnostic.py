"""Pure diagnostic: frozen input spectra, quantized steps, coherence, local tracking.

No model training. Rotation phase tests explicitly skipped when units lack evidence.
"""
import os
os.environ['OPENBLAS_NUM_THREADS']='2'
os.environ['OMP_NUM_THREADS']='2'
import time, json, hashlib, sys, zipfile
from pathlib import Path
import numpy as np
import pandas as pd
from scipy.signal import welch, coherence, find_peaks
from scipy.stats import spearmanr, pearsonr
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parent
PROJECT=ROOT.parents[1]
PREV=PROJECT/'analysis_outputs/cz_h1_light_whitening_20261006'
RAW=PROJECT/'实验数据1-张.xlsx'
ATTACH=Path('C:/Users/27437/.codex/attachments/b8cb230e-0862-43cd-82cb-aafe540f238a/已粘贴的文本.txt')
EXPECTED='c3428966fe006572809156ee5e3f488264b8206b19b20887dcd00840bb26fbc3'
START=time.perf_counter()
DT=2.;FS=.5;BAND=(.125,.155);REF=.138;SURROGATES=199
BREAKS={'Rod_1':[0,334,688,734], 'Rod_2':[0,745]}
CHANNELS={'crystal_rotation':'晶转速度','crucible_rotation':'埚转速度','main_heater_power':'主加热功率','crystal_lift':'晶升速度','crucible_lift':'埚升速度','heater_temperature':'加热元件温度'}

def sha(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def js(name,obj):(ROOT/name).write_text(json.dumps(obj,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
def log(*s):
    print(f'[{time.perf_counter()-START:.1f}s]',*s,flush=True)
    if time.time()-ATTACH.stat().st_ctime>1800:raise RuntimeError('STOP: task 30-minute budget exceeded')

def spectral(x):
    nperseg=min(1024,len(x));noverlap=nperseg//2
    f,p=welch(x,fs=FS,nperseg=nperseg,noverlap=noverlap,detrend='constant',scaling='density')
    return f,p,nperseg,1+(len(x)-nperseg)//(nperseg-noverlap)

def peak_summary(f,p):
    band=(f>=BAND[0])&(f<=BAND[1]);idx=np.flatnonzero(band)
    j=int(idx[np.argmax(p[band])]) if len(idx) else 0
    dom=int(np.argmax(p[1:])+1);nearest=int(np.argmin(abs(f-REF)))
    local,_=find_peaks(p)
    close=local[(f[local]>=.135)&(f[local]<=.142)]
    q=int(close[np.argmax(p[close])]) if len(close) else None
    background=(f>=.10)&(f<=.18)&(~band)
    return dict(dominant_frequency_Hz=float(f[dom]),dominant_period_seconds=float(1/f[dom]),band_power=float(p[band].sum()*(f[1]-f[0])),band_peak_frequency_Hz=float(f[j]),band_peak_PSD=float(p[j]),nearest_0138_frequency_Hz=float(f[nearest]),nearest_0138_PSD=float(p[nearest]),near_0138_local_peak_frequency_Hz=float(f[q]) if q is not None else np.nan,near_0138_peak_to_background=float(p[q]/np.median(p[background])) if q is not None and background.any() else np.nan,frequency_resolution_Hz=float(f[1]-f[0]))

def permutation_peak_p(x,f,p,rng):
    mask=(f>=BAND[0])&(f<=BAND[1]);observed=float(p[mask].max())
    null=[]
    for _ in range(SURROGATES):
        _,q,_,_=spectral(rng.permutation(x));null.append(float(q[mask].max()))
    return (1+int(np.sum(np.asarray(null)>=observed)))/(SURROGATES+1),float(np.quantile(null,.95))

def unit_audit():
    wb=load_workbook(RAW,read_only=False,data_only=False)
    evidence=[]
    for sheet in wb:
        headers=[cell.value for cell in sheet[1]]
        for name in ['晶转速度','埚转速度']:
            cell=sheet.cell(1,headers.index(name)+1)
            evidence.append(dict(sheet=sheet.title,field=name,header=cell.value,header_comment=cell.comment.text if cell.comment else None,number_format=cell.number_format,visible_state=sheet.sheet_state))
    props={k:getattr(wb.properties,k,None) for k in ['title','subject','description','keywords']}
    js('ROTATION_UNIT_EVIDENCE.json',dict(status='UNCONFIRMED',workbook_headers=evidence,descriptive_properties=props,search_log='UNIT_SEARCH.log',interpretation='No RPM/rev/s/rad/s evidence in CZ code, metadata, workbook headers/comments. Other benchmark rad/s matches are not CZ evidence.'))
    text='# 旋转速度单位只读审计\n\n结论：**UNCONFIRMED**。无法从当前项目证明晶转速度/埚转速度的物理单位。没有进行 rpm→Hz 换算、旋转相位积分或 phase model fitting。\n\n'
    text+='证据来源：\n\n- 原工作簿 Sheet1/Sheet2：字段只有“晶转速度”“埚转速度”；表头无单位注释，数值格式不含物理单位。工作簿描述属性没有单位说明，见 `ROTATION_UNIT_EVIDENCE.json`。\n- 原 H1 `cz_l256_nowcast.py`：仅将上述字段映射为 crystal_rotation/crucible_rotation，并声明 RAW_PERIOD_SECONDS=2.0；没有旋转单位声明。\n- `residual_rollout_authority_20260928/PRISM_INDUSTRIAL_BENCHMARK_V1/dataset_registry/cz_czochralski/` 和本项目 CZ data code：未发现 rpm、rev/min、r/min、rev/s、rad/s 或转/分证据。全项目可见文本中的 rad/s 命中来自 NanoDrone benchmark，与 CZ 无关，见 `UNIT_SEARCH.log`。\n- 既有 `analysis_outputs/cz_heater_identifiability_20261001/REPORT.md` 的数据说明同样明确工作簿没有可验证的单位说明。\n\n'
    text+='按用户要求，速度分布仅按原始记录单位报告（`ROTATION_SPEED_DISTRIBUTIONS.csv`）。数值大小与某种常见单位相似不构成单位证据。恒速旋转可具有相位相关输出但没有速度波动 coherence，因此未确认单位或低 speed coherence 均不能排除 H_R。\n'
    (ROOT/'ROTATION_UNIT_AUDIT.md').write_text(text,encoding='utf-8')
    wb.close()

def main():
    assert sha(RAW)==EXPECTED
    js('DIAGNOSTIC_PREREGISTRATION.json',dict(task='SEVEN_SECOND_PEAK_SOURCE_DIAGNOSTIC',classification='POST_HOC_DESCRIPTIVE_DIAGNOSTIC',raw_sha256=EXPECTED,sampling_seconds=DT,fs_Hz=FS,registered_band_Hz=BAND,reference_frequency_Hz=REF,Welch_nperseg=1024,Welch_noverlap=512,short_segments='nperseg=min(1024,n); noverlap=floor(nperseg/2); single-average coherence withheld',local_windows='nonoverlapping 2048 raw rows per segment, retain final remainder only if >=1536; Welch on 2047 derivatives',lattice_q=.01,lattice_tolerance=5e-5,spectral_null=dict(kind='within-segment random permutation, test maximum PSD over fixed registered band',draws=SURROGATES,seed=20261006,scope='long-segment signals only',limitation='Tests temporal ordering versus exchangeability, not physics; no global multiple-comparison adjustment'),coherence='scipy.signal.coherence with identical segment/Welch parameters; no phase inference; no independent significance claim',phase_condition='only if rotation units verified; otherwise explicit NOT_PERFORMED',training='none',input_artifact_hashes={n:sha(PREV/n) for n in ['POINTWISE_VALIDATION.parquet','POINTWISE_TEST.parquet','GLOBAL_SELECTION_FREEZE.json','WHITENING_METRICS.csv','SPECTRAL_SUMMARY.csv']}))
    unit_audit();log('unit audit UNCONFIRMED; phase tests NOT_PERFORMED')
    sheets=pd.read_excel(RAW,sheet_name=None)
    raw={f'Rod_{i+1}':sheets[f'Sheet{i+1}'] for i in range(2)}
    segments={}
    for rod,frame in raw.items():
        edges=BREAKS[rod]+[len(frame)]
        for j,(a,b) in enumerate(zip(edges[:-1],edges[1:])):
            segments[f'{rod}_segment_{j}']=dict(rod=rod,start=a,stop=b,frame=frame.iloc[a:b].reset_index(drop=True))
    rng=np.random.default_rng(20261006)
    quant=[];intervals=[];steps=[];output_psd=[];control_psd=[];control_summary=[];coh=[];coh_summary=[];tracking=[];speeds=[];lattice=[];alignment=[];paired=[];phase=[]
    plot_data={};rod_interval_store={rod:[] for rod in raw}
    for entity,seg in segments.items():
        frame=seg['frame'];rod=seg['rod'];D=frame['晶体直径'].to_numpy(float);delta=np.diff(D);events=np.flatnonzero(delta!=0);arrivals=np.diff(events)
        rod_interval_store[rod].append(arrivals)
        gap=abs(D-np.round(D/.01)*.01)
        float32_encoding=np.float32(np.round(D,2)).astype(float)
        lattice.append(dict(rod=rod,entity=entity,n=len(D),D_unique_exact=len(np.unique(D)),dD_unique_exact=len(np.unique(delta)),dD_unique_hundredth_grid=len(np.unique(np.diff(np.round(D/.01)))),dD_zero_ratio=float(np.mean(delta==0)),lattice_q=.01,max_lattice_distance=float(gap.max()),lattice_fraction_within_5e_5=float(np.mean(gap<=5e-5)),float32_grid_match_fraction=float(np.mean(abs(D-float32_encoding)<1e-9)),step_events=len(events),step_intervals=len(arrivals),interval_median_samples=float(np.median(arrivals)) if len(arrivals) else np.nan,interval_mode_samples=int(pd.Series(arrivals).mode().iloc[0]) if len(arrivals) else np.nan,interval_3_4_fraction=float(np.mean((arrivals>=3)&(arrivals<=4))) if len(arrivals) else np.nan,independent_geometric_3_4_fraction=float((1-np.mean(delta!=0))**2*np.mean(delta!=0)+(1-np.mean(delta!=0))**3*np.mean(delta!=0))))
        for interval,count in zip(*np.unique(arrivals,return_counts=True)):intervals.append(dict(rod=rod,entity=entity,scope='segment',interval_samples=int(interval),interval_seconds=int(interval)*2,count=int(count),fraction=float(count/len(arrivals))))
        magnitudes=np.abs(delta[delta!=0]);grid_steps=np.round(magnitudes/.01).astype(int)
        for g,c in zip(*np.unique(grid_steps,return_counts=True)):steps.append(dict(rod=rod,entity=entity,step_size_grid=.01*int(g),nonzero_count=int(c),nonzero_fraction=float(c/len(grid_steps))))
        signals={'D':D,'dD':delta,'nonzero':(delta!=0).astype(float),'absstep':abs(delta),'signstep':np.sign(delta),'signed_unit_step':.01*np.sign(delta)}
        for signal,x in signals.items():
            f,p,nper,K=spectral(x);q=peak_summary(f,p)
            q.update(rod=rod,entity=entity,scope='segment',signal=signal,n=len(x),nperseg=nper,Welch_averages=K)
            if len(x)>=1536 and signal not in ['signed_unit_step']:
                pvalue,null95=permutation_peak_p(x,f,p,rng)
                q.update(band_max_permutation_p=pvalue,band_max_permutation_null95=null95)
            else:q.update(band_max_permutation_p=np.nan,band_max_permutation_null95=np.nan)
            q['near_0138_prominent_local_peak']=bool(q['near_0138_peak_to_background']>=3.)
            quant.append(q)
            for fr,ps in zip(f,p):output_psd.append(dict(rod=rod,entity=entity,signal=signal,frequency_Hz=fr,PSD=ps,nperseg=nper,Welch_averages=K))
            if len(x)>=1536:plot_data[rod,signal]=(f,p)
        # Constant signed-unit-step removes amplitude variation while preserving signs and event times.
        f,pdD,_,_=spectral(delta);_,pu,_,_=spectral(.01*np.sign(delta));band=(f>=BAND[0])&(f<=BAND[1])
        paired.append(dict(rod=rod,entity=entity,scope='segment',dD_signed_unit_step_correlation=float(np.corrcoef(delta,.01*np.sign(delta))[0,1]),amplitude_only_residual_variance_fraction=float(np.var(delta-.01*np.sign(delta))/np.var(delta)),signed_unit_band_power_over_dD=float(pu[band].sum()/pdD[band].sum()),signed_unit_band_peak_Hz=float(f[np.flatnonzero(band)[np.argmax(pu[band])]])))
        for channel,name in CHANNELS.items():
            x=frame[name].to_numpy(float)
            speeds.append(dict(rod=rod,entity=entity,channel=channel,unit='UNCONFIRMED_RAW_UNIT',p10=float(np.quantile(x,.1)),median=float(np.median(x)),p90=float(np.quantile(x,.9)),min=float(x.min()),max=float(x.max())))
            f,p,nper,K=spectral(x);q=peak_summary(f,p);q.update(rod=rod,entity=entity,channel=channel,n=len(x),nperseg=nper,Welch_averages=K);control_summary.append(q)
            for fr,ps in zip(f,p):control_psd.append(dict(rod=rod,entity=entity,channel=channel,scope='raw_segment',frequency_Hz=fr,PSD=ps,nperseg=nper,Welch_averages=K))
            coherence_record(rod,entity,'raw_segment','dD','',channel,x[1:],delta,coh,coh_summary)
        for begin in range(0,len(frame),2048):
            stop=min(begin+2048,len(frame))
            if stop-begin<1536:continue
            g=frame.iloc[begin:stop];dd=np.diff(g['晶体直径'].to_numpy(float));f,p,nper,K=spectral(dd);q=peak_summary(f,p)
            tracking.append(dict(rod=rod,entity=entity,window_start_local=begin,window_stop_exclusive_local=stop,global_start=seg['start']+begin,global_stop_exclusive=seg['start']+stop,n=len(dd),peak_frequency_0125_0155=q['band_peak_frequency_Hz'],peak_power=q['band_peak_PSD'],median_crystal_rotation=float(g['晶转速度'].median()),median_crucible_rotation=float(g['埚转速度'].median()),nonzero_step_rate=float(np.mean(dd!=0)),nonzero_step_rate_Hz=float(np.mean(dd!=0)*FS),Welch_averages=K))
        log('raw segment processed',entity,len(frame))
    # Pool only static counts / within-segment inter-arrivals, never join derivative streams.
    for rod,frame in raw.items():
        ls=[v for v in lattice if v['rod']==rod];length=sum(v['n']-1 for v in ls)
        dD_unique=set();dd_grid_unique=set();nonzero_count=0
        for entity,seg in segments.items():
            if seg['rod']==rod:
                dd=np.diff(seg['frame']['晶体直径'].to_numpy(float));dD_unique.update(dd);dd_grid_unique.update(np.diff(np.round(seg['frame']['晶体直径'].to_numpy(float)/.01)));nonzero_count+=int(np.sum(dd!=0))
        arr=np.concatenate(rod_interval_store[rod]);event_rate=nonzero_count/length
        row=dict(rod=rod,entity='WITHIN_SEGMENT_STATIC_COUNT_POOL',n=len(frame),D_unique_exact=int(frame['晶体直径'].nunique()),dD_unique_exact=len(dD_unique),dD_unique_hundredth_grid=len(dd_grid_unique),dD_zero_ratio=1-event_rate,lattice_q=.01,max_lattice_distance=max(v['max_lattice_distance'] for v in ls),lattice_fraction_within_5e_5=sum(v['n']*v['lattice_fraction_within_5e_5'] for v in ls)/len(frame),float32_grid_match_fraction=sum(v['n']*v['float32_grid_match_fraction'] for v in ls)/len(frame),step_events=nonzero_count,step_intervals=len(arr),interval_median_samples=float(np.median(arr)),interval_mode_samples=int(pd.Series(arr).mode().iloc[0]),interval_3_4_fraction=float(np.mean((arr>=3)&(arr<=4))),independent_geometric_3_4_fraction=float(event_rate*(1-event_rate)**2+event_rate*(1-event_rate)**3))
        lattice.append(row)
        for gap,count in zip(*np.unique(arr,return_counts=True)):intervals.append(dict(rod=rod,entity='WITHIN_SEGMENT_STATIC_COUNT_POOL',scope='rod',interval_samples=int(gap),interval_seconds=int(gap)*2,count=int(count),fraction=float(count/len(arr))))
    for split in ['validation','test']:
        pt=pd.read_parquet(PREV/f'POINTWISE_{split.upper()}.parquet')
        for arm,g0 in pt.groupby('arm',sort=False):
            g=g0.sort_values('global_origin').reset_index(drop=True)
            # All delivered tails are single contiguous segments; verify every target exactly.
            assert g.entity.nunique()==1 and (g.origin.diff().dropna()==1).all()
            entity=g.entity.iloc[0];seg=segments[entity];rod=seg['rod'];local=g.origin.to_numpy(int);glob=g.global_origin.to_numpy(int)
            assert np.array_equal(glob,local+seg['start'])
            D=seg['frame']['晶体直径'].to_numpy(float);y=D[local]-D[local-1]
            assert np.max(abs(y-g.y_true))<1e-12
            assert np.max(abs(g.y_true-g.base_prediction-g.base_residual))<1e-12
            alignment.append(dict(rod=rod,arm=arm,split=split,entity=entity,n=len(g),global_origin_min=int(glob.min()),global_origin_max=int(glob.max()),raw_target_alignment_max_error=float(np.max(abs(y-g.y_true))),status='PASS'))
            for channel,name in CHANNELS.items():
                x=seg['frame'][name].to_numpy(float)[local]
                for target,col in [('dD','y_true'),('base_residual','base_residual')]:coherence_record(rod,entity,split,target,arm,channel,x,g[col].to_numpy(float),coh,coh_summary)
                # Input t-1 aligns to predictive-information convention; MSC differs from simultaneous phase.
                xlag=seg['frame'][name].to_numpy(float)[local-1]
                coherence_record(rod,entity,split,'base_residual',arm,channel,xlag,g.base_residual.to_numpy(float),coh,coh_summary,input_alignment='t-1')
            for target in ['dD','base_residual']:
                for phase_arm in ['crystal_phase','crucible_phase','both_phases']:
                    phase.append(dict(rod=rod,entity=entity,arm=arm,split=split,target=target,phase_model=phase_arm,status='NOT_PERFORMED_ROTATION_UNITS_UNCONFIRMED',heldout_RMSE_change=np.nan,E160_change=np.nan,E900_change=np.nan,band_power_change=np.nan,peak_0138_power_change=np.nan))
            # Re-report delivered residual correction, without fitting or changing a model.
            for signal,col in [('dD','y_true'),('nonzero',None),('absstep',None),('signstep',None)]:
                a=g.y_true.to_numpy(float);x=a if col else (a!=0).astype(float) if signal=='nonzero' else abs(a) if signal=='absstep' else np.sign(a)
                f,p,nper,K=spectral(x);q=peak_summary(f,p);q.update(rod=rod,entity=entity,scope=split,arm=arm,signal=signal,n=len(x),nperseg=nper,Welch_averages=K,band_max_permutation_p=np.nan,band_max_permutation_null95=np.nan,near_0138_prominent_local_peak=bool(q['near_0138_peak_to_background']>=3.));quant.append(q)
            for signal,column in [('base_residual','base_residual'),('final_residual','final_residual')]:
                f,p,nper,K=spectral(g[column].to_numpy());q=peak_summary(f,p);q.update(rod=rod,entity=entity,scope=split,arm=arm,signal=signal,n=len(g),nperseg=nper,Welch_averages=K);quant.append(q)
    track=pd.DataFrame(tracking);correlations=[];local_events=[]
    for rod,g in track.groupby('rod'):
        for variable in ['median_crystal_rotation','median_crucible_rotation','nonzero_step_rate_Hz']:
            pear=pearsonr(g[variable],g.peak_frequency_0125_0155);spear=spearmanr(g[variable],g.peak_frequency_0125_0155)
            drift=np.diff(g[variable].to_numpy());df=np.diff(g.peak_frequency_0125_0155.to_numpy())
            diffr=float(pearsonr(drift,df).statistic)
            tt=g.global_start.to_numpy(float)
            rxt=float(pearsonr(g[variable],tt).statistic);ryt=float(pearsonr(g.peak_frequency_0125_0155,tt).statistic)
            partial=(float(pear.statistic)-rxt*ryt)/np.sqrt((1-rxt*rxt)*(1-ryt*ryt))
            correlations.append(dict(rod=rod,variable=variable,n=len(g),Pearson_r=float(pear.statistic),Spearman_r=float(spear.statistic),first_difference_Pearson_r=diffr,linear_time_partial_Pearson_r=float(partial),direction_match_fraction=float(np.mean(np.sign(drift[df!=0])==np.sign(df[df!=0]))),warning='Descriptive only; time/regime confounding; dependent windows; no causal p-value claimed'))
        for signal in ['nonzero','absstep','signstep']:
            locations=[]
            for _,row in g.iterrows():
                frame=segments[row.entity]['frame'].iloc[int(row.window_start_local):int(row.window_stop_exclusive_local)]
                dd=np.diff(frame['晶体直径'].to_numpy(float));x=(dd!=0).astype(float) if signal=='nonzero' else abs(dd) if signal=='absstep' else np.sign(dd)
                f,p,_,_=spectral(x);q=peak_summary(f,p)
                locations.append(dict(rod=rod,entity=row.entity,global_start=int(row.global_start),signal=signal,**q))
            # Same registered band and resolution as dD; no shift to event-train dominant band.
            local_events.extend(locations)
    frames={'LOCAL_EVENT_PEAKS':local_events,'QUANTIZATION_DIAGNOSTIC':quant,'STEP_INTERVAL_HISTOGRAM':intervals,'STEP_SIZE_COUNTS':steps,'LATTICE_AND_INTERVAL_SUMMARY':lattice,'OUTPUT_PSD':output_psd,'CONTROL_PSD':control_psd,'CONTROL_SPECTRAL_SUMMARY':control_summary,'CONTROL_OUTPUT_COHERENCE':coh_summary,'COHERENCE_SPECTRA':coh,'ROTATION_SPEED_DISTRIBUTIONS':speeds,'ROTATION_PHASE_HELDOUT':phase,'LOCAL_PEAK_TRACKING':tracking,'LOCAL_TRACKING_CORRELATIONS':correlations,'SIGNED_UNIT_STEP_DIAGNOSTIC':paired,'POINTWISE_ALIGNMENT_AUDIT':alignment}
    for name,rows in frames.items():pd.DataFrame(rows).to_csv(ROOT/f'{name}.csv',index=False)
    make_plots(raw,track,plot_data,pd.DataFrame(intervals))
    js('RUN_SUMMARY.json',dict(status='COMPLETE',analysis_wall_seconds=time.perf_counter()-START,request_to_analysis_seconds=time.time()-ATTACH.stat().st_ctime,raw_sha256=sha(RAW),model_training_performed=False,rotation_units_confirmed=False,phase_models_fitted=0,prior_pointwise_alignment='PASS',prior_input_hashes_unchanged={name:sha(PREV/name)==h for name,h in json.loads((ROOT/'DIAGNOSTIC_PREREGISTRATION.json').read_text())['input_artifact_hashes'].items()}))
    log('analysis complete')

def coherence_record(rod,entity,scope,target,arm,channel,x,y,full,summary,input_alignment='t'):
    nper=min(1024,len(y));over=nper//2;K=1+(len(y)-nper)//(nper-over)
    if K<2:
        summary.append(dict(rod=rod,entity=entity,scope=scope,arm=arm,target=target,channel=channel,input_alignment=input_alignment,n=len(y),nperseg=nper,Welch_averages=K,status='NOT_IDENTIFIABLE_SINGLE_WELCH_AVERAGE',coherence_nearest_0138=np.nan,band_mean_coherence=np.nan,band_max_coherence=np.nan,band_max_frequency_Hz=np.nan));return
    f,c=coherence(x,y,fs=FS,nperseg=nper,noverlap=over,detrend='constant')
    mask=(f>=BAND[0])&(f<=BAND[1]);indices=np.flatnonzero(mask);j=int(indices[np.argmax(c[mask])]);near=int(np.argmin(abs(f-REF)))
    summary.append(dict(rod=rod,entity=entity,scope=scope,arm=arm,target=target,channel=channel,input_alignment=input_alignment,n=len(y),nperseg=nper,Welch_averages=K,status='LOW_AVERAGING_DESCRIPTIVE_ONLY' if K<5 else 'DESCRIPTIVE',nearest_frequency_Hz=float(f[near]),coherence_nearest_0138=float(c[near]),band_mean_coherence=float(c[mask].mean()),band_max_coherence=float(c[j]),band_max_frequency_Hz=float(f[j]),independent_segment_null_pointwise95=float(1-.05**(1/(K-1))),null_limit_note='Illustrative only: overlapping windows/dependence violate independent averages; band max not multiplicity-adjusted'))
    for fr,v in zip(f,c):full.append(dict(rod=rod,entity=entity,scope=scope,arm=arm,target=target,channel=channel,input_alignment=input_alignment,frequency_Hz=fr,coherence=v,nperseg=nper,Welch_averages=K))

def make_plots(raw,track,plot_data,hist):
    for rod in raw:
        label=rod.replace('_','')
        fig,axes=plt.subplots(2,2,figsize=(13,8))
        for signal,color in [('dD','#2676b8'),('nonzero','#c98528'),('absstep','#5e997c'),('signstep','#aa5a91')]:
            f,p=plot_data[rod,signal];mask=(f>=.10)&(f<=.18)
            axes[0,0].plot(f[mask],p[mask]/p[(f>=.125)&(f<=.155)].max(),label=signal,color=color)
            if signal in ['dD','absstep']:axes[0,1].semilogy(f,p,label=signal,color=color)
            else:axes[1,0].semilogy(f,p,label=signal,color=color)
        for ax in [axes[0,0],axes[0,1],axes[1,0]]:
            ax.axvspan(*BAND,color='gray',alpha=.12);ax.set_xlabel('Frequency (Hz)');ax.legend();ax.grid(alpha=.15)
        axes[0,0].set_title('Registered-band normalized PSD: long segment');axes[0,0].set_ylabel('PSD / own band maximum')
        axes[0,1].set_title('Diameter increments and absolute steps');axes[1,0].set_title('Binary event train and signed event train')
        g=hist[(hist.rod==rod)&(hist.scope=='rod')&(hist.interval_samples<=12)]
        axes[1,1].bar(g.interval_seconds,g.fraction,width=1.6,color='#59839c');axes[1,1].set_xlabel('Nonzero event interval (seconds)');axes[1,1].set_ylabel('Fraction');axes[1,1].axvspan(6,8,color='gray',alpha=.2);axes[1,1].set_title('Within-segment step-event inter-arrivals')
        fig.suptitle(f'{label}: quantized steps and ~7 s peak');fig.tight_layout();fig.savefig(ROOT/f'{label}_dD_step_event_PSD.png',dpi=160);plt.close(fig)
        g=track[track.rod==rod]
        fig,axes=plt.subplots(2,3,figsize=(14,8))
        timeh=g.global_start*DT/3600
        axes[0,0].plot(timeh,g.peak_frequency_0125_0155,'o-',color='#2676b8');axes[0,0].set_ylabel('Band peak (Hz)');axes[0,0].set_xlabel('Rod time (hours)');axes[0,0].set_ylim(*BAND)
        for j,variable in enumerate(['median_crystal_rotation','median_crucible_rotation','nonzero_step_rate_Hz']):
            ax=axes[1,j];scatter=ax.scatter(g[variable],g.peak_frequency_0125_0155,c=timeh,cmap='viridis',s=42);ax.set_xlabel(variable+' (raw units)' if j<2 else 'Event rate (Hz)');ax.set_ylabel('Band peak (Hz)');fig.colorbar(scatter,ax=ax,label='Rod time (hours)')
        axes[0,1].plot(timeh,g.median_crystal_rotation,'o-',label='crystal');axes[0,1].set_ylabel('Crystal speed (unit unconfirmed)');axes[0,1].set_xlabel('Rod time (hours)')
        axes[0,2].plot(timeh,g.median_crucible_rotation,'o-',label='crucible',color='#c98528');axes[0,2].set_ylabel('Crucible speed (unit unconfirmed)');axes[0,2].set_xlabel('Rod time (hours)')
        for ax in axes.flat:ax.grid(alpha=.15)
        fig.suptitle(f'{label}: local peak tracking | descriptive; time confounded | phase NOT TESTED');fig.tight_layout();fig.savefig(ROOT/f'{label}_rotation_vs_peak.png',dpi=160);plt.close(fig)

if __name__=='__main__':main()
